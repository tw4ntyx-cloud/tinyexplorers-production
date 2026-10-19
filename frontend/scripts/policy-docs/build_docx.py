"""Build the editable Word (.docx) version of a policy from the same data as the PDF.

Run build.mjs first (it writes scripts/policy-docs/out/<slug>.json), then from frontend/:

    python3 scripts/policy-docs/build_docx.py [slug ...]

Requires python-docx. Writes to docs/policies/docx/<docxFile from metadata.json>.
"""
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

HERE = Path(__file__).resolve().parent
FRONTEND = HERE.parent.parent
# Editable copies live with the policy sources, not on the public website.
OUT = FRONTEND.parent / "docs" / "policies" / "docx"

INK = RGBColor(0x1F, 0x1F, 0x1F)
MUTED = RGBColor(0x5B, 0x57, 0x53)
ORANGE = RGBColor(0xE8, 0x62, 0x2A)
SERIF = "Georgia"
SANS = "Arial"
LABELLED = re.compile(r"^([A-Z][A-Za-z /&-]{2,40}): (.+)$")
QUOTED = re.compile(r'^(".*?")\s*(.*)$', re.S)


def run(paragraph, text, size=10.5, bold=False, italic=False, color=INK, font=SANS):
    r = paragraph.add_run(text)
    r.font.name = font
    r._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    return r


def para(doc, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.35
    return p


def bottom_border(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("val", "single"), ("sz", "4"), ("space", "4"), ("color", "E2DDD6")):
        bottom.set(qn(f"w:{key}"), value)
    borders.append(bottom)
    p_pr.append(borders)


def build(slug):
    data = json.loads((HERE / "out" / f"{slug}.json").read_text())
    policy, meta = data["policy"], data["meta"]
    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Pt(61)
        header = section.header.paragraphs[0]
        run(header, "Tiny Explorers Nursery & Preschool", 10, font=SERIF)
        run(header, "\nWhere Learning Blossoms, and Adventures Begin", 8, italic=True, color=MUTED, font=SERIF)
        footer = section.footer.paragraphs[0]
        run(footer, "tinyexplorersbda.com  ·  © 2026 Tiny Explorers Nursery & Preschool", 7.5, color=MUTED)

    # Document control table (same fields as the PDF header)
    fields = [("Version", meta["version"]), ("Doc ID", meta["docId"]), ("Effective", meta["effective"]),
              ("Review", meta["review"]), ("Approved by", meta["approvedBy"])]
    table = doc.add_table(rows=2, cols=len(fields))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (label, value) in enumerate(fields):
        run(table.cell(0, i).paragraphs[0], label.upper(), 7, color=MUTED)
        run(table.cell(1, i).paragraphs[0], value, 9)

    label = "POLICY DOCUMENT" if meta.get("pdf", True) else "POLICY DOCUMENT · EDITABLE COPY OF THE PUBLISHED PDF"
    run(para(doc, 2), label, 8, color=ORANGE)
    title = para(doc, 4)
    run(title, policy["title"], 24, font=SERIF)
    summary = para(doc, 10)
    run(summary, policy["summary"], 10.5, color=MUTED)
    bottom_border(summary)

    sections = [dict(s, blocks=list(s["blocks"])) for s in policy["content"]]
    if sections and sections[0]["blocks"] and sections[0]["blocks"][0]["type"] == "paragraph":
        run(para(doc, 10), sections[0]["blocks"].pop(0)["text"], 11.5)

    for index, section in enumerate(sections, start=1):
        number = para(doc, 0)
        number.paragraph_format.keep_with_next = True
        run(number, f"{index:02d}", 8, color=ORANGE)
        heading = para(doc, 6)
        heading.paragraph_format.keep_with_next = True
        run(heading, section["title"], 16, font=SERIF)
        bottom_border(heading)
        for block in section["blocks"]:
            if block["type"] == "paragraph":
                p = para(doc)
                m = LABELLED.match(block["text"])
                if m:
                    run(p, f"{m.group(1)}: ", bold=True)
                    run(p, m.group(2))
                else:
                    run(p, block["text"])
            elif block["type"] == "list":
                for item in block["items"]:
                    p = para(doc, 2)
                    p.paragraph_format.left_indent = Pt(14)
                    run(p, "— ", color=ORANGE)
                    run(p, item)
            elif block["type"] == "note":
                run(para(doc, 2), (block.get("title") or "Note").upper(), 7.5, color=ORANGE)
                m = QUOTED.match(block["text"])
                if m:
                    run(para(doc, 2), m.group(1), 12, italic=True, font=SERIF)
                    run(para(doc), m.group(2))
                else:
                    run(para(doc), block["text"])
            else:
                raise ValueError(f"Unknown block type: {block['type']}")

    doc.add_page_break()
    run(para(doc, 6), "REVISION HISTORY", 8, color=MUTED)
    history = doc.add_table(rows=1, cols=4)
    history.style = "Table Grid"
    for i, label in enumerate(("Version", "Date", "Summary", "Author")):
        run(history.cell(0, i).paragraphs[0], label.upper(), 7, color=MUTED)
    for rev in meta["revisions"]:
        cells = history.add_row().cells
        for i, key in enumerate(("version", "date", "summary", "author")):
            run(cells[i].paragraphs[0], rev[key], 9)
    run(para(doc, 2), "\nAPPROVED BY", 7.5, color=MUTED)
    run(para(doc, 0), meta["approvedBy"])
    run(para(doc, 12), "Tiny Explorers Nursery & Preschool", 8, color=MUTED)
    run(para(doc), "DATE: ______________________", 7.5, color=MUTED)

    doc.core_properties.title = policy["title"]
    doc.core_properties.author = "Tiny Explorers Nursery & Preschool"
    doc.core_properties.subject = f"{meta['docId']} v{meta['version']}"
    OUT.mkdir(parents=True, exist_ok=True)
    target = OUT / meta["docxFile"]
    doc.save(target)
    print(f"DOCX {target.name} ({target.stat().st_size} bytes)")


if __name__ == "__main__":
    slugs = sys.argv[1:] or [p.stem for p in (HERE / "out").glob("*.json")]
    for s in slugs:
        build(s)
